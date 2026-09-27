"""Safe text extraction from uploaded company documents.

Nothing in a document is executed: PDF and DOCX are only parsed for text, macros,
scripts and external links are ignored, and the content must match its extension
before a parser sees it.
"""

from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass

MEDIA_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain",
    "md": "text/markdown",
}
# What browsers and curl may declare for each format; anything else is refused.
ACCEPTED_DECLARED_TYPES = {
    "pdf": {"application/pdf", "application/x-pdf"},
    "docx": {MEDIA_TYPES["docx"], "application/zip"},
    "txt": {"text/plain"},
    "md": {"text/markdown", "text/x-markdown", "text/plain"},
}
GENERIC_DECLARED_TYPES = {"", "application/octet-stream", "binary/octet-stream"}

MAX_FILENAME_LENGTH = 120
_SAFE_NAME_CHARS = re.compile(r"[^\w .()\-]", re.UNICODE)
_CONTROL = re.compile(r"[\u0000-\u0008\u000b\u000c\u000e-\u001f\u007f​-‏  ⁠﻿]")


class ExtractionError(Exception):
    """A document that cannot become knowledge; `message` is safe to show an admin."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class ExtractedDocument:
    format: str
    media_type: str
    text: str
    pages: int | None = None


def file_format(filename: str) -> str | None:
    """Supported format by extension, or None. `.docm` and friends are not DOCX."""
    _, dot, extension = filename.rpartition(".")
    extension = extension.lower() if dot else ""
    return extension if extension in MEDIA_TYPES else None


def media_type_for(fmt: str) -> str:
    return MEDIA_TYPES[fmt]


def declared_type_matches(fmt: str, declared: str | None) -> bool:
    value = (declared or "").split(";", 1)[0].strip().lower()
    return value in GENERIC_DECLARED_TYPES or value in ACCEPTED_DECLARED_TYPES[fmt]


def safe_filename(raw: str) -> str:
    """Display name only: no directories, control or markup characters. Never used as a path."""
    name = unicodedata.normalize("NFC", raw or "")
    name = re.split(r"[\\/]", name)[-1]
    name = _SAFE_NAME_CHARS.sub("", _CONTROL.sub("", name))
    name = re.sub(r"\s+", " ", name).strip(" .")
    if len(name) > MAX_FILENAME_LENGTH:
        stem, dot, extension = name.rpartition(".")
        keep = MAX_FILENAME_LENGTH - len(extension) - 1
        name = f"{stem[:keep]}.{extension}" if dot and 0 < len(extension) <= 10 else name[:MAX_FILENAME_LENGTH]
    return name or "document"


def normalize_text(text: str) -> str:
    """NFC, LF newlines, no control characters, single spaces, at most one blank line."""
    text = unicodedata.normalize("NFC", text).replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL.sub("", text).replace("\t", " ").replace(" ", " ")
    lines = [re.sub(r" {2,}", " ", line).strip() for line in text.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def extract_document(filename: str, content: bytes, *, max_chars: int, max_pdf_pages: int) -> ExtractedDocument:
    fmt = file_format(filename)
    if fmt is None:
        raise ExtractionError("unsupported_format", "Поддерживаются только PDF, DOCX, TXT и MD.")
    _check_signature(fmt, content)
    pages = None
    if fmt == "pdf":
        raw, pages = _pdf_text(content, max_pdf_pages)
    elif fmt == "docx":
        raw = _docx_text(content)
    else:
        raw = _plain_text(content)
    text = normalize_text(raw)
    if not re.search(r"\w", text):
        raise ExtractionError("no_text", "В документе нет текста, который можно прочитать. Сканы без распознавания не поддерживаются.")
    if len(text) > max_chars:
        raise ExtractionError("too_large", f"В документе больше {max_chars} символов текста. Разделите его на части.")
    return ExtractedDocument(format=fmt, media_type=MEDIA_TYPES[fmt], text=text, pages=pages)


# --- signatures -----------------------------------------------------------------

def _check_signature(fmt: str, content: bytes) -> None:
    head = content[:1024].lstrip()
    is_pdf = head.startswith(b"%PDF-")
    is_zip = content.startswith(b"PK\x03\x04")
    if fmt == "pdf" and not is_pdf:
        raise _mismatch()
    if fmt == "docx" and not is_zip:
        raise _mismatch()
    if fmt in {"txt", "md"}:
        if is_pdf or is_zip or _looks_binary(content):
            raise _mismatch()


def _looks_binary(content: bytes) -> bool:
    sample = content[:8192]
    if b"\x00" in sample:
        return True
    if sample.startswith((b"MZ", b"\x7fELF", b"\x89PNG", b"\xff\xd8\xff", b"GIF8", b"\xd0\xcf\x11\xe0")):
        return True
    controls = sum(1 for byte in sample if byte < 9 or 13 < byte < 32)
    return bool(sample) and controls / len(sample) > 0.05


def _mismatch() -> ExtractionError:
    return ExtractionError("content_mismatch", "Содержимое файла не соответствует его расширению.")


# --- parsers ----------------------------------------------------------------------

def _plain_text(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ExtractionError("not_utf8", "Текстовый файл должен быть в кодировке UTF-8.") from error


def _pdf_text(content: bytes, max_pages: int) -> tuple[str, int]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(content), strict=False)
        if reader.is_encrypted:
            raise ExtractionError("encrypted", "PDF защищён паролем. Загрузите версию без шифрования.")
        page_count = len(reader.pages)
        if page_count > max_pages:
            raise ExtractionError("too_many_pages", f"В PDF больше {max_pages} страниц.")
        # Only the text layer is read: annotations, JavaScript, forms and links are ignored.
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    except ExtractionError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError, AttributeError, IndexError, RecursionError) as error:
        raise ExtractionError("corrupt", "PDF повреждён, прочитать его не получилось.") from error
    except Exception as error:  # pypdf raises many internal types on hostile input
        raise ExtractionError("corrupt", "PDF повреждён, прочитать его не получилось.") from error
    return text, page_count


def _docx_text(content: bytes) -> str:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = set(archive.namelist())
            if "word/document.xml" not in names or "[Content_Types].xml" not in names:
                raise ExtractionError("corrupt", "Файл не похож на документ Word.")
            if any(name.lower().endswith("vbaproject.bin") for name in names):
                raise ExtractionError("content_mismatch", "Документы с макросами не принимаются.")
        document = Document(io.BytesIO(content))
    except ExtractionError:
        raise
    except (zipfile.BadZipFile, PackageNotFoundError, KeyError, ValueError) as error:
        raise ExtractionError("corrupt", "Документ Word повреждён, прочитать его не получилось.") from error
    except Exception as error:
        raise ExtractionError("corrupt", "Документ Word повреждён, прочитать его не получилось.") from error

    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        level = _heading_level(paragraph.style.name if paragraph.style is not None else "")
        parts.append(f"{'#' * level} {text}" if level else text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _heading_level(style_name: str) -> int:
    if style_name == "Title":
        return 1
    match = re.fullmatch(r"(?:Heading|Заголовок) (\d)", style_name or "")
    return min(int(match.group(1)), 6) if match else 0
