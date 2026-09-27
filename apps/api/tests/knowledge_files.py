"""Real test documents built in memory: no binary fixtures in the repository."""

from __future__ import annotations

import io
import zipfile

from docx import Document
from pypdf import PdfWriter


def _escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def pdf_with_text(*lines: str) -> bytes:
    """A minimal valid one-page PDF whose text pypdf can extract (Latin-1 Helvetica)."""
    commands = ["BT", "/F1 12 Tf", "72 720 Td", "14 TL"]
    for line in lines:
        commands.append(f"({_escape(line)}) Tj T*")
    commands.append("ET")
    stream = "\n".join(commands).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def blank_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def encrypted_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt(user_password="secret", owner_password="owner")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def docx_with(heading: str, paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = Document()
    document.add_heading(heading, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for column_index, value in enumerate(row):
                grid.cell(row_index, column_index).text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def zip_without_docx() -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("readme.txt", "not a word document")
    return out.getvalue()
